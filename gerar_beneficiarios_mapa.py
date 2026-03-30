from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.styles import Font

def main():
	setores = [1787,
2058,
2057,
2077,
2524,
5152,
2492,
2441,
2710,
1951,
1232,
1957,
2076,
1918,
5234,
1849,
2001,
2073,
1391,
5247,
5233,
2026,
2074,
1950,
2085,
2050,
2075,
1857,
5249,
1755,
2885,
2439,
5189,
5213,
1848,
]
	pessoas = []

	planilha = load_workbook("quadroGeral.xlsx")
	funcionarios = planilha.active
	colunas = funcionarios.max_row + 1

	for row in range(2, colunas):

		setor = funcionarios[f"AC{row}"].value

		if setor in setores:

			pessoa = {
			"matricula": funcionarios[f"C{row}"].value,
			"nome": funcionarios[f"D{row}"].value,
			"ug": funcionarios[f"AC{row}"].value,
			"operacao": funcionarios[f"AD{row}"].value,
			"cargo": funcionarios[f"F{row}"].value,
			"site": funcionarios[f"A{row}"].value,
			"matricula_tt": funcionarios[f"AA{row}"].value
			}

			pessoas.append(pessoa)


	gerarBeneficiariosMapa(pessoas)



def gerarBeneficiariosMapa(pessoas):

	beneficiarios = Workbook()
	aba = beneficiarios.active
	aba.title

	fundo = PatternFill(start_color = "002060", end_color = "002060", fill_type = "solid")
	fonte = Font(color = "FFFFFF")

	aba["A1"].value = "MATRICULA"
	aba["A1"].fill = fundo
	aba["A1"].font = fonte

	aba["B1"].value = "NOME"
	aba["B1"].fill = fundo
	aba["B1"].font = fonte

	aba["C1"].value = "UG"
	aba["C1"].fill = fundo
	aba["C1"].font = fonte

	aba["D1"].value = "OPERAÇÃO"
	aba["D1"].fill = fundo
	aba["D1"].font = fonte

	aba["E1"].value = "CARGO"
	aba["E1"].fill = fundo
	aba["E1"].font = fonte

	aba["F1"].value = "SITE"
	aba["F1"].fill = fundo
	aba["F1"].font = fonte

	aba["G1"].value = "MATRICULA TT"
	aba["G1"].fill = fundo
	aba["G1"].font = fonte

	linha = 2

	for pessoa in pessoas:
		aba[f"A{linha}"] = pessoa["matricula"]
		aba[f"B{linha}"] = pessoa["nome"]
		aba[f"C{linha}"] = pessoa["ug"]
		aba[f"D{linha}"] = pessoa["operacao"]
		aba[f"E{linha}"] = pessoa["cargo"]
		aba[f"F{linha}"] = pessoa["site"]
		aba[f"G{linha}"] = pessoa["matricula_tt"]


		linha += 1

	beneficiarios.save("TESTE.xlsx")


if __name__ == '__main__':
	main()
